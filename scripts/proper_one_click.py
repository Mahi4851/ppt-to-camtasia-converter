from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import sys
import time
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


TICKS_PER_SECOND = 705_600_000
CANVAS_WIDTH = 1920
CANVAS_HEIGHT = 1080
DEFAULT_SLIDE_SECONDS = 5.0
STANDARD_FONT_NAME = "Open Sans"
STANDARD_FONT_SIZE_PT = 34.0
DEFAULT_TRANSITION_NAME = "Fade"
DEFAULT_TRANSITION_SECONDS = 0.5


def ensure_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def save_json(data: Any, path: Path) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")


def die(message: str) -> None:
    print("")
    print("ERROR: " + message)
    raise SystemExit(1)


def sec_to_ticks(seconds: float) -> int:
    return int(round(float(seconds) * TICKS_PER_SECOND))


def emu_to_px(value: Any, total_emu: int, total_px: int) -> float:
    if value is None or not total_emu:
        return 0.0
    return round((float(value) / float(total_emu)) * float(total_px), 3)


def safe_text(value: Any) -> str:
    return str(value or "").replace("\r", "\n").strip()


def clean_name(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9_. -]+", "_", value.strip())
    return re.sub(r"\s+", " ", value).strip() or "Untitled"


def find_input_pptx(input_dir: Path) -> Path:
    files = [
        p for p in input_dir.glob("*.pptx")
        if not p.name.startswith("~$") and p.is_file()
    ]
    if not files:
        die(f"No .pptx file found in input folder: {input_dir}")
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0]


def import_pptx_modules():
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except Exception as exc:
        die("python-pptx is missing. Run run_all.bat so requirements install first. Details: " + repr(exc))
    return Presentation, MSO_SHAPE_TYPE


def iter_shapes(shapes: Any, shape_type_group: Any, prefix: str = ""):
    for shape in shapes:
        path = f"{prefix}/{shape.shape_id}:{shape.name}" if prefix else f"{shape.shape_id}:{shape.name}"
        yield shape, path
        if shape.shape_type == shape_type_group:
            yield from iter_shapes(shape.shapes, shape_type_group, path)


def iter_shapes_with_slide_geometry(
    shapes: Any,
    shape_type_group: Any,
    prefix: str = "",
    x_offset: float = 0.0,
    y_offset: float = 0.0,
    x_scale: float = 1.0,
    y_scale: float = 1.0,
    parent_rotation: float = 0.0,
    z_path: tuple[int, ...] = (),
):
    """Yield shapes with coordinates resolved into the original slide coordinate space.

    PowerPoint group children use the group's child coordinate system (chOff/chExt).
    Treating those values as slide coordinates can shift or scale grouped artwork.
    This iterator preserves the existing extraction behaviour while resolving that
    group transform before X/Y/width/height are converted to pixels.
    """
    for index, shape in enumerate(shapes, start=1):
        path = f"{prefix}/{shape.shape_id}:{shape.name}" if prefix else f"{shape.shape_id}:{shape.name}"
        left = float(getattr(shape, "left", 0) or 0)
        top = float(getattr(shape, "top", 0) or 0)
        width = float(getattr(shape, "width", 0) or 0)
        height = float(getattr(shape, "height", 0) or 0)
        local_rotation = float(getattr(shape, "rotation", 0) or 0)
        current_z_path = z_path + (index,)
        geometry = {
            "left_emu": x_offset + x_scale * left,
            "top_emu": y_offset + y_scale * top,
            "width_emu": abs(x_scale * width),
            "height_emu": abs(y_scale * height),
            "rotation": (parent_rotation + local_rotation) % 360.0,
            "z_order_path": list(current_z_path),
            "group_depth": max(len(current_z_path) - 1, 0),
            "parent_group_path": prefix,
        }
        yield shape, path, geometry

        if shape.shape_type == shape_type_group:
            try:
                xfrm = shape.element.grpSpPr.xfrm
                ch_off_x = float(xfrm.chOff.x)
                ch_off_y = float(xfrm.chOff.y)
                ch_ext_x = float(xfrm.chExt.cx)
                ch_ext_y = float(xfrm.chExt.cy)
            except Exception:
                ch_off_x = left
                ch_off_y = top
                ch_ext_x = width or 1.0
                ch_ext_y = height or 1.0

            child_x_scale = x_scale * (width / ch_ext_x if ch_ext_x else 1.0)
            child_y_scale = y_scale * (height / ch_ext_y if ch_ext_y else 1.0)
            child_x_offset = x_offset + x_scale * left - child_x_scale * ch_off_x
            child_y_offset = y_offset + y_scale * top - child_y_scale * ch_off_y
            yield from iter_shapes_with_slide_geometry(
                shape.shapes,
                shape_type_group,
                path,
                child_x_offset,
                child_y_offset,
                child_x_scale,
                child_y_scale,
                parent_rotation + local_rotation,
                current_z_path,
            )


def paragraph_text(paragraph: Any) -> str:
    return "".join(run.text for run in paragraph.runs).strip()


def first_text_run(paragraphs: list[Any]) -> Any | None:
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if safe_text(run.text):
                return run
    return None


def color_to_hex(color: Any) -> str | None:
    if color is None:
        return None
    try:
        rgb = color.rgb
        if rgb is not None:
            return "#" + str(rgb)
    except Exception:
        return None
    return None


def points(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return round(float(value.pt), 3)
    except Exception:
        return None


def bool_or_none(value: Any) -> bool | None:
    if value is None:
        return None
    return bool(value)


def alignment_name(paragraph: Any) -> str:
    alignment = getattr(paragraph, "alignment", None)
    if alignment is None:
        return "left"
    try:
        name = alignment.name.lower()
    except Exception:
        name = str(alignment).lower()
    if "center" in name:
        return "center"
    if "right" in name:
        return "right"
    return "left"


def shape_position(shape: Any, slide_w: int, slide_h: int) -> dict[str, float]:
    return {
        "x_px": emu_to_px(getattr(shape, "left", 0), slide_w, CANVAS_WIDTH),
        "y_px": emu_to_px(getattr(shape, "top", 0), slide_h, CANVAS_HEIGHT),
        "width_px": max(emu_to_px(getattr(shape, "width", 0), slide_w, CANVAS_WIDTH), 0.01),
        "height_px": max(emu_to_px(getattr(shape, "height", 0), slide_h, CANVAS_HEIGHT), 0.01),
        "rotation": float(getattr(shape, "rotation", 0) or 0),
        "x_emu": float(getattr(shape, "left", 0) or 0),
        "y_emu": float(getattr(shape, "top", 0) or 0),
        "width_emu": float(getattr(shape, "width", 0) or 0),
        "height_emu": float(getattr(shape, "height", 0) or 0),
    }


def resolved_shape_position(geometry: dict[str, Any], slide_w: int, slide_h: int) -> dict[str, Any]:
    left = float(geometry.get("left_emu") or 0.0)
    top = float(geometry.get("top_emu") or 0.0)
    width = float(geometry.get("width_emu") or 0.0)
    height = float(geometry.get("height_emu") or 0.0)
    return {
        "x_px": emu_to_px(left, slide_w, CANVAS_WIDTH),
        "y_px": emu_to_px(top, slide_h, CANVAS_HEIGHT),
        "width_px": max(emu_to_px(width, slide_w, CANVAS_WIDTH), 0.01),
        "height_px": max(emu_to_px(height, slide_h, CANVAS_HEIGHT), 0.01),
        "rotation": float(geometry.get("rotation") or 0.0),
        "x_emu": left,
        "y_emu": top,
        "width_emu": width,
        "height_emu": height,
        "z_order_path": list(geometry.get("z_order_path") or []),
        "group_depth": int(geometry.get("group_depth") or 0),
        "parent_group_path": str(geometry.get("parent_group_path") or ""),
    }


def style_from_paragraphs(paragraphs: list[Any]) -> dict[str, Any]:
    paragraph = next((p for p in paragraphs if paragraph_text(p)), paragraphs[0] if paragraphs else None)
    run = first_text_run(paragraphs)
    font = run.font if run is not None else getattr(paragraph, "font", None)
    return {
        "font_family": getattr(font, "name", None) if font is not None else None,
        "font_size_pt": points(getattr(font, "size", None)) if font is not None else None,
        "font_color": color_to_hex(getattr(font, "color", None)) if font is not None else None,
        "bold": bool_or_none(getattr(font, "bold", None)) if font is not None else None,
        "italic": bool_or_none(getattr(font, "italic", None)) if font is not None else None,
        "underline": bool_or_none(getattr(font, "underline", None)) if font is not None else None,
        "alignment": alignment_name(paragraph) if paragraph is not None else "left",
    }


def slide_duration_seconds(slide: Any) -> float:
    try:
        transitions = slide.element.xpath(".//p:transition")
        for transition in transitions:
            adv_tm = transition.get("advTm")
            if adv_tm is not None:
                return max(float(adv_tm) / 1000.0, 0.25)
    except Exception:
        pass
    return DEFAULT_SLIDE_SECONDS


PPTX_NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def parse_millis(value: Any) -> float | None:
    text = str(value or "").strip()
    if not text or text == "indefinite":
        return None
    try:
        return max(float(text), 0.0)
    except Exception:
        return None


def ctn_delay_ms(ctn: ET.Element) -> float:
    delays: list[float] = []
    for cond in ctn.findall("./p:stCondLst/p:cond", PPTX_NS):
        parsed = parse_millis(cond.get("delay"))
        if parsed is not None:
            delays.append(parsed)
    return max(delays) if delays else 0.0


def action_duration_ms(action: ET.Element) -> float:
    ctn = action.find(".//p:cBhvr/p:cTn", PPTX_NS)
    if ctn is None:
        return 0.0
    parsed = parse_millis(ctn.get("dur"))
    return parsed if parsed is not None else 0.0


def action_target_shape_ids(action: ET.Element) -> list[int]:
    ids: list[int] = []
    for target in action.findall(".//p:spTgt", PPTX_NS):
        raw = target.get("spid")
        try:
            ids.append(int(raw))  # type: ignore[arg-type]
        except Exception:
            continue
    return ids


def animation_effect_name(action: ET.Element, ctn_attrs: dict[str, Any]) -> str:
    name = action.get("filter") or action.get("transition") or ""
    if not name and ctn_attrs.get("presetClass"):
        name = str(ctn_attrs.get("presetClass"))
    return name or local_name(action.tag)


def extract_animation_items(source_pptx: Path, slide_count: int) -> list[dict[str, Any]]:
    animation_items: list[dict[str, Any]] = []
    action_tags = {"animEffect", "anim", "animMotion", "animScale", "animRot", "cmd", "set"}

    try:
        archive = zipfile.ZipFile(source_pptx)
    except Exception:
        return animation_items

    with archive:
        slide_xml_names: list[tuple[int, str]] = []
        for name in archive.namelist():
            match = re.fullmatch(r"ppt/slides/slide(\d+)\.xml", name)
            if match:
                slide_xml_names.append((int(match.group(1)), name))
        slide_xml_names.sort()

        for slide_number, name in slide_xml_names:
            if slide_number > slide_count:
                continue
            try:
                root = ET.fromstring(archive.read(name))
            except Exception:
                continue
            timing = root.find(".//p:timing", PPTX_NS)
            if timing is None:
                continue

            order = 0

            def walk(node: ET.Element, inherited_ms: float, ctn_attrs: dict[str, Any]) -> None:
                nonlocal order
                current_ms = inherited_ms
                current_attrs = dict(ctn_attrs)
                if local_name(node.tag) == "cTn":
                    current_ms += ctn_delay_ms(node)
                    for attr in ("presetClass", "presetID", "presetSubtype", "nodeType"):
                        if node.get(attr) is not None:
                            current_attrs[attr] = node.get(attr)

                node_name = local_name(node.tag)
                if node_name in action_tags:
                    target_ids = action_target_shape_ids(node)
                    if target_ids:
                        order += 1
                        duration_ms = action_duration_ms(node)
                        for shape_id in target_ids:
                            animation_items.append({
                                "slide_number": slide_number,
                                "shape_id": shape_id,
                                "order": order,
                                "action": node_name,
                                "animation_class": current_attrs.get("presetClass") or "",
                                "effect": animation_effect_name(node, current_attrs),
                                "start_time_sec": round(current_ms / 1000.0, 6),
                                "effect_duration_sec": round(duration_ms / 1000.0, 6),
                                "preset_id": current_attrs.get("presetID") or "",
                                "preset_subtype": current_attrs.get("presetSubtype") or "",
                                "node_type": current_attrs.get("nodeType") or "",
                            })

                for child in list(node):
                    walk(child, current_ms, current_attrs)

            walk(timing, 0.0, {})

    return animation_items


def build_shape_animation_timing(
    animation_items: list[dict[str, Any]],
    slide_timing: list[dict[str, Any]],
) -> dict[tuple[int, int], dict[str, Any]]:
    slide_duration = {int(item["slide_number"]): float(item["duration_sec"]) for item in slide_timing}
    by_shape: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for item in animation_items:
        by_shape.setdefault((int(item["slide_number"]), int(item["shape_id"])), []).append(item)

    result: dict[tuple[int, int], dict[str, Any]] = {}
    for key, events in by_shape.items():
        slide_number, _shape_id = key
        events.sort(key=lambda event: (float(event.get("start_time_sec") or 0.0), int(event.get("order") or 0)))
        entrance_effects = [event for event in events if event.get("animation_class") == "entr" and event.get("action") != "set"]
        visibility_sets = [event for event in events if event.get("action") == "set"]
        exits = [event for event in events if event.get("animation_class") == "exit"]
        emph_or_motion = [event for event in events if event.get("animation_class") in {"emph", "path"} or event.get("action") in {"animMotion", "animScale", "animRot"}]

        first = events[0]
        first_entrance = entrance_effects[0] if entrance_effects else (visibility_sets[0] if visibility_sets else None)
        first_exit = exits[0] if exits else None
        if first_entrance is not None:
            start_sec = float(first_entrance.get("start_time_sec") or 0.0)
            primary = first_entrance
        else:
            start_sec = 0.0
            primary = emph_or_motion[0] if emph_or_motion else first

        slide_len = slide_duration.get(slide_number, DEFAULT_SLIDE_SECONDS)
        if first_exit is not None:
            end_sec = min(
                float(first_exit.get("start_time_sec") or slide_len) + float(first_exit.get("effect_duration_sec") or 0.0),
                slide_len,
            )
        else:
            end_sec = slide_len
        if end_sec <= start_sec:
            end_sec = min(slide_len, start_sec + 0.1)

        result[key] = {
            "start_time_sec": round(start_sec, 6),
            "duration_sec": round(max(end_sec - start_sec, 0.1), 6),
            "end_time_sec": round(end_sec, 6),
            "animation_order": int(primary.get("order") or 0),
            "animation_effect": primary.get("effect") or "",
            "animation_action": primary.get("action") or "",
            "animation_class": primary.get("animation_class") or "",
            "effect_duration_sec": primary.get("effect_duration_sec") or 0.0,
            "all_animation_events": events,
        }
    return result


def apply_animation_timing_to_items(extracted: "ExtractedData") -> None:
    timing_by_shape = build_shape_animation_timing(extracted.animation_items, extracted.slide_timing)

    for item in extracted.shape_items:
        key = (int(item["slide_number"]), int(item["shape_id"]))
        timing = timing_by_shape.get(key)
        if timing:
            item.update(timing)
        else:
            slide = extracted.slide_timing[int(item["slide_number"]) - 1]
            item.setdefault("start_time_sec", 0.0)
            item.setdefault("duration_sec", float(slide["duration_sec"]))
            item.setdefault("end_time_sec", float(slide["duration_sec"]))

    paragraph_entrances: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for event in extracted.animation_items:
        if event.get("animation_class") == "entr" and event.get("action") != "set":
            paragraph_entrances.setdefault((int(event["slide_number"]), int(event["shape_id"])), []).append(event)
    for events in paragraph_entrances.values():
        events.sort(key=lambda event: (float(event.get("start_time_sec") or 0.0), int(event.get("order") or 0)))

    text_by_shape: dict[tuple[int, int], list[dict[str, Any]]] = {}
    for item in extracted.text_items:
        text_by_shape.setdefault((int(item["slide_number"]), int(item["shape_id"])), []).append(item)
    for key, items in text_by_shape.items():
        items.sort(key=lambda item: int(item.get("paragraph_index") or 1))
        paragraph_events = paragraph_entrances.get(key, [])
        shared_timing = timing_by_shape.get(key)
        for index, item in enumerate(items):
            if paragraph_events and len(paragraph_events) == len(items):
                event = paragraph_events[index]
                slide = extracted.slide_timing[int(item["slide_number"]) - 1]
                start_sec = float(event.get("start_time_sec") or 0.0)
                end_sec = float(slide["duration_sec"])
                item.update({
                    "start_time_sec": round(start_sec, 6),
                    "duration_sec": round(max(end_sec - start_sec, 0.1), 6),
                    "end_time_sec": round(end_sec, 6),
                    "animation_order": int(event.get("order") or 0),
                    "animation_effect": event.get("effect") or "",
                    "animation_action": event.get("action") or "",
                    "animation_class": event.get("animation_class") or "",
                    "effect_duration_sec": event.get("effect_duration_sec") or 0.0,
                })
            elif shared_timing:
                item.update({k: v for k, v in shared_timing.items() if k != "all_animation_events"})


@dataclass
class ExtractedData:
    slide_timing: list[dict[str, Any]]
    text_items: list[dict[str, Any]]
    shape_items: list[dict[str, Any]]
    animation_items: list[dict[str, Any]]
    slide_count: int
    slide_width_emu: int
    slide_height_emu: int


def shape_type_name(shape: Any) -> str:
    try:
        return shape.shape_type.name
    except Exception:
        return str(getattr(shape, "shape_type", "unknown"))


def is_text_only_shape(shape: Any) -> bool:
    has_text = getattr(shape, "has_text_frame", False) and safe_text(shape.text_frame.text)
    has_table = getattr(shape, "has_table", False)
    try:
        if not has_text or has_table or shape_type_name(shape) not in {"TEXT_BOX", "PLACEHOLDER"}:
            return False
        return not shape_fill_hex(shape) and not shape_line_hex(shape)
    except Exception:
        return bool(has_text and not has_table)


def is_group_container_shape(shape: Any) -> bool:
    return shape_type_name(shape) == "GROUP"


def simple_native_kind(shape: Any) -> str | None:
    if getattr(shape, "has_table", False):
        return None
    shape_type = shape_type_name(shape)
    if shape_type not in {"AUTO_SHAPE", "LINE"}:
        return None
    try:
        auto_shape_type = getattr(shape, "auto_shape_type", None)
        name = auto_shape_type.name.lower() if auto_shape_type is not None else ""
    except Exception:
        name = ""
    shape_name = (getattr(shape, "name", "") or "").lower()
    combined = f"{name} {shape_name}"
    if "round" in combined and "rect" in combined:
        return "rounded_rectangle"
    if "rect" in combined:
        return "rectangle"
    if "oval" in combined or "ellipse" in combined:
        return "oval"
    if "line" in combined:
        return "line"
    return None


def shape_fill_hex(shape: Any) -> str | None:
    try:
        fill = shape.fill
        if fill and fill.fore_color and fill.fore_color.rgb:
            return "#" + str(fill.fore_color.rgb)
    except Exception:
        return None
    return None


def shape_line_hex(shape: Any) -> str | None:
    try:
        line = shape.line
        if line and line.color and line.color.rgb:
            return "#" + str(line.color.rgb)
    except Exception:
        return None
    return None


def extract_text_and_timing(source_pptx: Path, working_dir: Path) -> ExtractedData:
    Presentation, MSO_SHAPE_TYPE = import_pptx_modules()
    prs = Presentation(str(source_pptx))
    slide_w = int(prs.slide_width)
    slide_h = int(prs.slide_height)
    animation_items = extract_animation_items(source_pptx, len(prs.slides))
    text_items: list[dict[str, Any]] = []
    shape_items: list[dict[str, Any]] = []
    slide_timing: list[dict[str, Any]] = []
    current = 0.0

    for slide_number, slide in enumerate(prs.slides, start=1):
        duration = slide_duration_seconds(slide)
        slide_animation_events = [item for item in animation_items if int(item["slide_number"]) == slide_number]
        if slide_animation_events:
            animation_end = max(
                float(item.get("start_time_sec") or 0.0) + float(item.get("effect_duration_sec") or 0.0)
                for item in slide_animation_events
            )
            duration = max(duration, animation_end + 0.25)
        start = round(current, 6)
        end = round(current + duration, 6)
        title = f"Slide {slide_number:02d}"
        timing_xml = False
        try:
            timing_xml = bool(slide.element.xpath(".//p:timing"))
        except Exception:
            timing_xml = False

        slide_timing.append({
            "slide_number": slide_number,
            "title": title,
            "start_time_sec": start,
            "duration_sec": round(duration, 6),
            "end_time_sec": end,
            "animation_timing_detected": timing_xml,
            "animation_count": len(slide_animation_events),
        })

        for shape, shape_path, geometry in iter_shapes_with_slide_geometry(slide.shapes, MSO_SHAPE_TYPE.GROUP):
            position = resolved_shape_position(geometry, slide_w, slide_h)
            has_visible_text = getattr(shape, "has_text_frame", False) and safe_text(shape.text_frame.text)
            has_table = getattr(shape, "has_table", False)
            native_kind = simple_native_kind(shape)
            item_shape_type = shape_type_name(shape)
            if not is_text_only_shape(shape) and not is_group_container_shape(shape):
                shape_items.append({
                    "id": f"S{slide_number:03d}_SHAPE{shape.shape_id}",
                    "slide_number": slide_number,
                    "shape_id": shape.shape_id,
                    "shape_name": shape.name,
                    "shape_path": shape_path,
                    "shape_type": item_shape_type,
                    "asset_kind": "picture" if item_shape_type == "PICTURE" else "shape",
                    "native_kind": native_kind,
                    "classification": "separate_png_layer",
                    "editable_candidate": "native_camtasia_shape" if native_kind else "",
                    "has_text": bool(has_visible_text),
                    "has_table": bool(has_table),
                    "fill_color": shape_fill_hex(shape),
                    "line_color": shape_line_hex(shape),
                    "start_time_sec": 0.0,
                    "duration_sec": duration,
                    "end_time_sec": duration,
                    **position,
                })

            if getattr(shape, "has_text_frame", False):
                paragraphs = [p for p in shape.text_frame.paragraphs if paragraph_text(p)]
                if not paragraphs:
                    continue
                base = resolved_shape_position(geometry, slide_w, slide_h)
                if paragraphs:
                    each_h = max(base["height_px"] / len(paragraphs), 12.0)
                for index, paragraph in enumerate(paragraphs, start=1):
                    text = paragraph_text(paragraph)
                    if not text:
                        continue
                    item_pos = dict(base)
                    item_pos["y_px"] = round(base["y_px"] + (index - 1) * each_h, 3)
                    item_pos["height_px"] = each_h
                    style = style_from_paragraphs([paragraph])
                    if index == 1:
                        maybe_title = text.replace("\n", " ").strip()
                        if maybe_title:
                            title = clean_name(maybe_title[:60])
                            slide_timing[-1]["title"] = title
                    text_items.append({
                        "id": f"S{slide_number:03d}_ID{shape.shape_id}_P{index:02d}",
                        "slide_number": slide_number,
                        "shape_id": shape.shape_id,
                        "shape_name": shape.name,
                        "shape_path": shape_path,
                        "text": text,
                        "text_kind": "paragraph",
                        "paragraph_index": index,
                        **item_pos,
                        **style,
                        "start_time_sec": 0.0,
                        "duration_sec": duration,
                        "end_time_sec": duration,
                    })

            if getattr(shape, "has_table", False):
                base = resolved_shape_position(geometry, slide_w, slide_h)
                rows = list(shape.table.rows)
                cols = list(shape.table.columns)
                cell_w = base["width_px"] / max(len(cols), 1)
                cell_h = base["height_px"] / max(len(rows), 1)
                for r_idx, row in enumerate(rows, start=1):
                    for c_idx, cell in enumerate(row.cells, start=1):
                        text = safe_text(cell.text)
                        if not text:
                            continue
                        text_items.append({
                            "id": f"S{slide_number:03d}_TABLE{shape.shape_id}_R{r_idx}_C{c_idx}",
                            "slide_number": slide_number,
                            "shape_id": shape.shape_id,
                            "shape_name": shape.name,
                            "shape_path": shape_path,
                            "text": text,
                            "text_kind": "table_cell",
                            "paragraph_index": 1,
                            "x_px": round(base["x_px"] + (c_idx - 1) * cell_w, 3),
                            "y_px": round(base["y_px"] + (r_idx - 1) * cell_h, 3),
                            "width_px": round(cell_w, 3),
                            "height_px": round(cell_h, 3),
                            "rotation": 0.0,
                            "font_family": None,
                            "font_size_pt": None,
                            "font_color": None,
                            "bold": None,
                            "italic": None,
                            "underline": None,
                            "alignment": "left",
                            "start_time_sec": 0.0,
                            "duration_sec": duration,
                            "end_time_sec": duration,
                        })
        current = end

    result = ExtractedData(
        slide_timing=slide_timing,
        text_items=text_items,
        shape_items=shape_items,
        animation_items=animation_items,
        slide_count=len(prs.slides),
        slide_width_emu=slide_w,
        slide_height_emu=slide_h,
    )
    apply_animation_timing_to_items(result)
    combined_visual_items = visual_manifest(result)
    save_json({
        "source_pptx": str(source_pptx),
        "canvas": {"width": CANVAS_WIDTH, "height": CANVAS_HEIGHT},
        "slide_size_emu": {"width": slide_w, "height": slide_h},
        "slide_timing": slide_timing,
        "text_items": text_items,
        "shape_items": shape_items,
        "visual_items": combined_visual_items,
        "animation_items": animation_items,
    }, working_dir / "extracted_text_and_timing.json")
    return result


def duplicate_without_text(source_pptx: Path, output_pptx: Path, report_path: Path) -> None:
    Presentation, MSO_SHAPE_TYPE = import_pptx_modules()
    shutil.copy2(source_pptx, output_pptx)
    prs = Presentation(str(output_pptx))
    report = {"cleared_text_shapes": [], "cleared_table_cells": []}
    for slide_number, slide in enumerate(prs.slides, start=1):
        for shape, shape_path in iter_shapes(slide.shapes, MSO_SHAPE_TYPE.GROUP):
            if getattr(shape, "has_text_frame", False):
                old = safe_text(shape.text_frame.text)
                if old:
                    shape.text_frame.clear()
                    report["cleared_text_shapes"].append({
                        "slide_number": slide_number,
                        "shape_id": shape.shape_id,
                        "shape_name": shape.name,
                        "shape_path": shape_path,
                        "text": old,
                    })
            if getattr(shape, "has_table", False):
                for r_idx, row in enumerate(shape.table.rows, start=1):
                    for c_idx, cell in enumerate(row.cells, start=1):
                        old = safe_text(cell.text)
                        if old:
                            cell.text = ""
                            report["cleared_table_cells"].append({
                                "slide_number": slide_number,
                                "shape_id": shape.shape_id,
                                "row": r_idx,
                                "column": c_idx,
                                "text": old,
                            })
    prs.save(str(output_pptx))
    save_json(report, report_path)


def duplicate_background_only(source_pptx: Path, output_pptx: Path, report_path: Path) -> None:
    Presentation, _MSO_SHAPE_TYPE = import_pptx_modules()
    shutil.copy2(source_pptx, output_pptx)
    prs = Presentation(str(output_pptx))
    report = {"removed_shapes": []}
    for slide_number, slide in enumerate(prs.slides, start=1):
        for shape in list(slide.shapes):
            report["removed_shapes"].append({
                "slide_number": slide_number,
                "shape_id": shape.shape_id,
                "shape_name": shape.name,
                "shape_type": shape_type_name(shape),
            })
            parent = shape.element.getparent()
            parent.remove(shape.element)
    prs.save(str(output_pptx))
    save_json(report, report_path)


def path_ids(shape_path: str) -> list[int]:
    ids: list[int] = []
    for part in shape_path.split("/"):
        raw = part.split(":", 1)[0]
        try:
            ids.append(int(raw))
        except Exception:
            pass
    return ids


def find_com_shape(container: Any, ids: list[int]) -> Any | None:
    current_container = container
    current_shape = None
    for wanted in ids:
        found = None
        try:
            count = int(current_container.Count)
        except Exception:
            return None
        for index in range(1, count + 1):
            shape = current_container.Item(index)
            try:
                if int(shape.Id) == wanted:
                    found = shape
                    break
            except Exception:
                continue
        if found is None:
            return None
        current_shape = found
        if wanted != ids[-1]:
            try:
                current_container = found.GroupItems
            except Exception:
                return None
    return current_shape


def find_com_shape_deep(container: Any, wanted_id: int, wanted_name: str | None = None) -> Any | None:
    try:
        count = int(container.Count)
    except Exception:
        return None
    for index in range(1, count + 1):
        shape = container.Item(index)
        try:
            if int(shape.Id) == wanted_id:
                return shape
        except Exception:
            pass
        try:
            if wanted_name and str(shape.Name) == wanted_name:
                return shape
        except Exception:
            pass
        try:
            nested = find_com_shape_deep(shape.GroupItems, wanted_id, wanted_name)
            if nested is not None:
                return nested
        except Exception:
            pass
    return None


def export_shape_pngs_with_powerpoint(no_text_pptx: Path, shapes_dir: Path, shape_items: list[dict[str, Any]]) -> dict[str, str]:
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        die("pywin32 / PowerPoint COM is missing. Microsoft PowerPoint is required for shape export. Details: " + repr(exc))

    ensure_dir(shapes_dir)
    by_slide: dict[int, list[dict[str, Any]]] = {}
    for item in shape_items:
        if item.get("classification") != "separate_png_layer":
            continue
        by_slide.setdefault(int(item["slide_number"]), []).append(item)

    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = 1
    exported: dict[str, str] = {}
    failures: list[dict[str, Any]] = []
    try:
        deck = app.Presentations.Open(str(no_text_pptx), WithWindow=False)
        try:
            for slide_number, items in by_slide.items():
                slide = deck.Slides(slide_number)
                for item in items:
                    shape = find_com_shape(slide.Shapes, path_ids(item.get("shape_path", "")))
                    lookup_strategy = "path"
                    if shape is None:
                        shape = find_com_shape_deep(
                            slide.Shapes,
                            int(item.get("shape_id") or 0),
                            str(item.get("shape_name") or ""),
                        )
                        lookup_strategy = "deep_id_or_name"
                    if shape is None:
                        failures.append({
                            "id": item["id"],
                            "slide_number": item.get("slide_number"),
                            "shape_id": item.get("shape_id"),
                            "shape_name": item.get("shape_name"),
                            "shape_path": item.get("shape_path"),
                            "reason": "shape_not_found",
                        })
                        continue
                    output = shapes_dir / f"{item['id']}_{clean_name(item.get('shape_name', 'shape'))}.png"
                    try:
                        shape.Export(str(output), 2)
                    except Exception:
                        try:
                            shape.Copy()
                            temp = slide.Shapes.PasteSpecial(2).Item(1)
                            temp.Export(str(output), 2)
                            temp.Delete()
                        except Exception as exc:
                            failures.append({
                                "id": item["id"],
                                "slide_number": item.get("slide_number"),
                                "shape_id": item.get("shape_id"),
                                "shape_name": item.get("shape_name"),
                                "shape_path": item.get("shape_path"),
                                "lookup_strategy": lookup_strategy,
                                "reason": repr(exc),
                            })
                            continue
                    if output.exists() and output.stat().st_size > 0:
                        exported[item["id"]] = str(output)
                    else:
                        failures.append({
                            "id": item["id"],
                            "slide_number": item.get("slide_number"),
                            "shape_id": item.get("shape_id"),
                            "shape_name": item.get("shape_name"),
                            "shape_path": item.get("shape_path"),
                            "lookup_strategy": lookup_strategy,
                            "reason": "empty_export",
                        })
        finally:
            deck.Close()
    finally:
        app.Quit()
    if failures:
        save_json({"failures": failures}, shapes_dir.parent.parent / "reports" / "shape_export_failures.json")
    return exported


def export_slide_pngs_with_powerpoint(no_text_pptx: Path, backgrounds_dir: Path, slide_count: int) -> list[Path]:
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        die("pywin32 / PowerPoint COM is missing. Microsoft PowerPoint is required for background export. Details: " + repr(exc))

    ensure_dir(backgrounds_dir)
    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = 1
    exported: list[Path] = []
    try:
        deck = app.Presentations.Open(str(no_text_pptx), WithWindow=False)
        try:
            actual_count = int(deck.Slides.Count)
            if actual_count != slide_count:
                die(f"Slide count changed after text removal. Expected {slide_count}, got {actual_count}.")
            for slide_number in range(1, actual_count + 1):
                output = backgrounds_dir / f"S{slide_number:03d}_Background.png"
                deck.Slides(slide_number).Export(str(output), "PNG", CANVAS_WIDTH, CANVAS_HEIGHT)
                if not output.exists() or output.stat().st_size <= 0:
                    die(f"PowerPoint did not export slide {slide_number} background.")
                exported.append(output)
        finally:
            deck.Close()
    finally:
        app.Quit()
    return exported


def export_backgrounds_without_slide_shapes_with_powerpoint(
    source_pptx: Path,
    backgrounds_dir: Path,
    working_dir: Path,
    slide_count: int,
) -> list[Path]:
    try:
        import win32com.client  # type: ignore
    except Exception as exc:
        die("pywin32 / PowerPoint COM is missing. Microsoft PowerPoint is required for background export. Details: " + repr(exc))

    ensure_dir(backgrounds_dir)
    temp_pptx = working_dir / "background_export_powerpoint_copy.pptx"
    shutil.copy2(source_pptx, temp_pptx)
    report = {"method": "powerpoint_delete_slide_shapes_then_export", "removed_shapes": []}
    app = win32com.client.DispatchEx("PowerPoint.Application")
    app.Visible = 1
    try:
        try:
            app.DisplayAlerts = 0
        except Exception:
            pass
        deck = app.Presentations.Open(str(temp_pptx), WithWindow=False)
        try:
            actual_count = int(deck.Slides.Count)
            if actual_count != slide_count:
                raise RuntimeError(f"Slide count changed during background export. Expected {slide_count}, got {actual_count}.")
            exported: list[Path] = []
            for slide_number in range(1, actual_count + 1):
                slide = deck.Slides(slide_number)
                for index in range(int(slide.Shapes.Count), 0, -1):
                    shape = slide.Shapes.Item(index)
                    try:
                        report["removed_shapes"].append({
                            "slide_number": slide_number,
                            "shape_id": int(shape.Id),
                            "shape_name": str(shape.Name),
                            "shape_type": str(shape.Type),
                        })
                    except Exception:
                        report["removed_shapes"].append({"slide_number": slide_number, "shape_id": "", "shape_name": "", "shape_type": ""})
                    shape.Delete()
                output = backgrounds_dir / f"S{slide_number:03d}_Background.png"
                slide.Export(str(output), "PNG", CANVAS_WIDTH, CANVAS_HEIGHT)
                if not output.exists() or output.stat().st_size <= 0:
                    raise RuntimeError(f"PowerPoint did not export clean background for slide {slide_number}.")
                exported.append(output)
            save_json(report, working_dir / "background_powerpoint_export_report.json")
            return exported
        finally:
            try:
                deck.Saved = True
            except Exception:
                pass
            deck.Close()
    finally:
        app.Quit()


def export_backgrounds_safely(
    source_pptx: Path,
    no_text_pptx: Path,
    background_only_pptx: Path,
    backgrounds_dir: Path,
    working_dir: Path,
    slide_count: int,
) -> tuple[list[Path], str]:
    try:
        backgrounds = export_backgrounds_without_slide_shapes_with_powerpoint(source_pptx, backgrounds_dir, working_dir, slide_count)
        return backgrounds, "background_without_slide_shapes"
    except Exception as exc:
        save_json({
            "status": "PRIMARY_BACKGROUND_EXPORT_FAILED",
            "reason": repr(exc),
            "message": "PowerPoint shape-deletion background export failed; trying generated background-only PPTX next.",
        }, working_dir / "background_powerpoint_export_failed.json")

    try:
        duplicate_background_only(source_pptx, background_only_pptx, working_dir / "background_only_report.json")
        backgrounds = export_slide_pngs_with_powerpoint(background_only_pptx, backgrounds_dir, slide_count)
        return backgrounds, "background_only"
    except Exception as exc:
        fallback_report = {
            "status": "FALLBACK_USED",
            "reason": repr(exc),
            "message": (
                "PowerPoint could not open the generated background-only PPTX. "
                "The runner exported backgrounds from the no-text PPTX instead, so conversion can continue."
            ),
            "impact": (
                "Some non-text shapes may appear both in the background PNG and as separate editable/deletable tracks. "
                "If duplicated, hide/delete the separate shape track or keep it as the editable object layer."
            ),
        }
        save_json(fallback_report, working_dir / "background_export_fallback_report.json")
        print("")
        print("WARNING: background-only PPTX could not be opened by PowerPoint.")
        print("Continuing with no-text slide backgrounds instead.")
        print("Details saved to working\\background_export_fallback_report.json")
        backgrounds = export_slide_pngs_with_powerpoint(no_text_pptx, backgrounds_dir, slide_count)
        return backgrounds, "no_text_fallback"


def hex_to_float_rgb(value: str | None) -> tuple[float, float, float]:
    if not value:
        return 0.25, 0.25, 0.25
    text = value.strip().lstrip("#")
    if len(text) != 6:
        return 0.25, 0.25, 0.25
    return int(text[0:2], 16) / 255.0, int(text[2:4], 16) / 255.0, int(text[4:6], 16) / 255.0


def camtasia_color(value: str | None) -> str:
    if not value:
        return "(64,64,64,255)"
    text = value.strip().lstrip("#")
    if len(text) != 6:
        return "(64,64,64,255)"
    return f"({int(text[0:2], 16)},{int(text[2:4], 16)},{int(text[4:6], 16)},255)"


def png_size(path: Path) -> tuple[int, int]:
    try:
        header = path.read_bytes()[:24]
        if header.startswith(b"\x89PNG\r\n\x1a\n"):
            return int.from_bytes(header[16:20], "big"), int.from_bytes(header[20:24], "big")
    except Exception:
        pass
    return CANVAS_WIDTH, CANVAS_HEIGHT


def make_source(source_id: int, path: Path, project_root: Path) -> dict[str, Any]:
    rel = "./" + path.relative_to(project_root).as_posix()
    width, height = png_size(path)
    return {
        "id": source_id,
        "src": rel,
        "rect": [0, 0, width, height],
        "lastMod": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y%m%dT%H%M%S"),
        "loudnessNormalization": True,
        "sourceTracks": [{
            "range": [0, 1],
            "type": 1,
            "editRate": 10_000_000,
            "trackRect": [0, 0, width, height],
            "sampleRate": 0,
            "bitDepth": 32,
            "numChannels": 0,
            "integratedLUFS": 100.0,
            "peakLevel": -1.0,
            "metaData": path.name + ";",
            "parameters": {},
        }],
        "metadata": {"timeAdded": datetime.now().strftime("%Y%m%dT%H%M%S.%f")},
    }


def track_wrapper(
    track_index: int,
    ident: str,
    medias: list[dict[str, Any]],
    transitions: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "trackIndex": track_index,
        "medias": medias,
        "transitions": transitions or [],
        "metadata": {"WinTrackHeight": "56", "IsLocked": "False"},
        "ident": ident,
        "audioMuted": False,
        "videoHidden": False,
        "magnetic": False,
        "matte": 0,
        "solo": False,
        "isAudioTrack": False,
    }


def item_clip_timing(item: dict[str, Any], fallback_duration_ticks: int) -> tuple[int, int]:
    fallback_duration_sec = max(float(fallback_duration_ticks) / float(TICKS_PER_SECOND), 0.1)
    start_sec = max(float(item.get("start_time_sec") or 0.0), 0.0)
    duration_sec = float(item.get("duration_sec") or max(fallback_duration_sec - start_sec, 0.1))
    if duration_sec <= 0:
        duration_sec = 0.1
    return sec_to_ticks(start_sec), max(sec_to_ticks(duration_sec), sec_to_ticks(0.1))


def fade_transitions_for_clip(clip: dict[str, Any]) -> list[dict[str, Any]]:
    duration = int(clip.get("duration") or 0)
    if duration <= 0:
        return []
    transition_duration = min(sec_to_ticks(DEFAULT_TRANSITION_SECONDS), max(duration // 3, sec_to_ticks(0.1)))
    attrs = {
        "bypass": False,
        "reverse": False,
        "trivial": False,
        "useAudioPreRoll": True,
        "useVisualPreRoll": True,
    }
    return [
        {
            "name": DEFAULT_TRANSITION_NAME,
            "duration": transition_duration,
            "rightMedia": int(clip["id"]),
            "attributes": dict(attrs),
        },
        {
            "name": DEFAULT_TRANSITION_NAME,
            "duration": transition_duration,
            "leftMedia": int(clip["id"]),
            "attributes": dict(attrs),
        },
    ]


def background_clip(clip_id: int, source_id: int, duration: int, slide_number: int) -> dict[str, Any]:
    return {
        "id": clip_id,
        "_type": "IMFile",
        "src": source_id,
        "trackNumber": 0,
        "trimStartSum": 0,
        "attributes": {"ident": f"S{slide_number:03d}_background"},
        "parameters": {
            "translation0": 0.0,
            "translation1": 0.0,
            "scale0": 1.0,
            "scale1": 1.0,
            "geometryCrop0": 0.0,
            "geometryCrop1": 0.0,
            "geometryCrop2": 0.0,
            "geometryCrop3": 0.0,
        },
        "effects": [],
        "start": 0,
        "duration": duration,
        "mediaStart": 0,
        "mediaDuration": 1,
        "scalar": 1,
        "metadata": {
            "audiateLinkedSession": "",
            "clipSpeedAttribute": {"type": "bool", "value": False},
            "colorAttribute": {"type": "color", "value": [0, 0, 0, 0]},
            "default-scale0": {"type": "double", "value": 1.0},
            "default-scale1": {"type": "double", "value": 1.0},
            "default-translation0": {"type": "double", "value": 0.0},
            "default-translation1": {"type": "double", "value": 0.0},
            "powerPointElementType": "background",
            "powerPointShapeName": "Background",
            "transitionName": DEFAULT_TRANSITION_NAME,
        },
        "animationTracks": {},
    }


def image_shape_clip(clip_id: int, source_id: int, item: dict[str, Any], image_path: Path, duration: int) -> dict[str, Any]:
    img_w, img_h = png_size(image_path)
    target_w = max(float(item.get("width_px") or img_w), 1.0)
    target_h = max(float(item.get("height_px") or img_h), 1.0)
    x = float(item.get("x_px") or 0.0)
    y = float(item.get("y_px") or 0.0)
    scale0 = target_w / max(float(img_w), 1.0)
    scale1 = target_h / max(float(img_h), 1.0)
    translation0 = x + target_w / 2.0 - CANVAS_WIDTH / 2.0
    translation1 = -(y + target_h / 2.0 - CANVAS_HEIGHT / 2.0)
    start_ticks, duration_ticks = item_clip_timing(item, duration)
    return {
        "id": clip_id,
        "_type": "IMFile",
        "src": source_id,
        "trackNumber": 0,
        "trimStartSum": 0,
        "attributes": {"ident": item["id"]},
        "parameters": {
            "translation0": translation0,
            "translation1": translation1,
            "scale0": scale0,
            "scale1": scale1,
            "geometryCrop0": 0.0,
            "geometryCrop1": 0.0,
            "geometryCrop2": 0.0,
            "geometryCrop3": 0.0,
        },
        "effects": [],
        "start": start_ticks,
        "duration": duration_ticks,
        "mediaStart": 0,
        "mediaDuration": 1,
        "scalar": 1,
        "metadata": {
            "audiateLinkedSession": "",
            "clipSpeedAttribute": {"type": "bool", "value": False},
            "default-scale0": {"type": "double", "value": scale0},
            "default-scale1": {"type": "double", "value": scale1},
            "default-translation0": {"type": "double", "value": translation0},
            "default-translation1": {"type": "double", "value": translation1},
            "powerPointElementType": "separate_shape_png",
            "powerPointShapeId": {"type": "int", "value": int(item.get("shape_id") or 0)},
            "powerPointShapeName": item.get("shape_name") or "",
            "powerPointAssetKind": item.get("asset_kind") or "shape",
            "powerPointReadingOrder": {"type": "int", "value": int(item.get("reading_order") or 0)},
            "powerPointRelationshipId": item.get("relationship_id") or "",
            "powerPointRelationshipReason": item.get("relationship_reason") or "",
            "powerPointRotation": {"type": "double", "value": float(item.get("rotation") or 0.0)},
            "powerPointZOrderPath": ".".join(str(v) for v in item.get("z_order_path") or []),
            "powerPointAnimationOrder": {"type": "int", "value": int(item.get("animation_order") or 0)},
            "powerPointAnimationEffect": item.get("animation_effect") or "",
            "powerPointAnimationClass": item.get("animation_class") or "",
            "transitionName": DEFAULT_TRANSITION_NAME,
        },
        "animationTracks": {},
    }


def native_shape_clip(clip_id: int, item: dict[str, Any], duration: int) -> dict[str, Any]:
    kind = item.get("native_kind") or "rectangle"
    shape_name = {
        "rounded_rectangle": "roundedRectangle",
        "rectangle": "rectangle",
        "oval": "ellipse",
        "line": "line",
    }.get(kind, "rectangle")
    width = max(float(item.get("width_px") or 100), 8.0)
    height = max(float(item.get("height_px") or 40), 8.0)
    x = float(item.get("x_px") or 0.0)
    y = float(item.get("y_px") or 0.0)
    translation0 = x + width / 2.0 - CANVAS_WIDTH / 2.0
    translation1 = -(y + height / 2.0 - CANVAS_HEIGHT / 2.0)
    fill = camtasia_color(item.get("fill_color"))
    outline = camtasia_color(item.get("line_color"))
    red, green, blue = hex_to_float_rgb(item.get("fill_color"))
    start_ticks, duration_ticks = item_clip_timing(item, duration)
    return {
        "id": clip_id,
        "_type": "Callout",
        "attributes": {"ident": item["id"], "autoRotateText": False},
        "def": {
            "kind": "remix",
            "shape": shape_name,
            "style": "basic",
            "height": height,
            "width": width,
            "fill": fill,
            "outline": outline,
            "outline-width": 1.0,
            "font": {
                "color-blue": blue,
                "color-green": green,
                "color-red": red,
                "color-opacity": 1.0,
                "size": STANDARD_FONT_SIZE_PT,
                "tracking": 0.0,
                "name": STANDARD_FONT_NAME,
                "weight": "Regular",
            },
        },
        "parameters": {
            "translation0": translation0,
            "translation1": translation1,
            "scale0": 1.0,
            "scale1": 1.0,
            "geometryCrop0": 0.0,
            "geometryCrop1": 0.0,
            "geometryCrop2": 0.0,
            "geometryCrop3": 0.0,
        },
        "effects": [],
        "start": start_ticks,
        "duration": duration_ticks,
        "mediaStart": 0,
        "mediaDuration": duration_ticks,
        "scalar": 1,
        "metadata": {
            "audiateLinkedSession": "",
            "clipSpeedAttribute": {"type": "bool", "value": False},
            "default-height": {"type": "double", "value": height},
            "default-width": {"type": "double", "value": width},
            "default-scale0": {"type": "double", "value": 1.0},
            "default-scale1": {"type": "double", "value": 1.0},
            "default-translation0": {"type": "double", "value": translation0},
            "default-translation1": {"type": "double", "value": translation1},
            "powerPointElementType": "native_shape",
            "powerPointNativeKind": kind,
            "powerPointShapeId": {"type": "int", "value": int(item.get("shape_id") or 0)},
            "powerPointShapeName": item.get("shape_name") or "",
            "powerPointReadingOrder": {"type": "int", "value": int(item.get("reading_order") or 0)},
            "powerPointRelationshipId": item.get("relationship_id") or "",
            "powerPointRelationshipReason": item.get("relationship_reason") or "",
            "powerPointRotation": {"type": "double", "value": float(item.get("rotation") or 0.0)},
            "powerPointZOrderPath": ".".join(str(v) for v in item.get("z_order_path") or []),
            "powerPointAnimationOrder": {"type": "int", "value": int(item.get("animation_order") or 0)},
            "powerPointAnimationEffect": item.get("animation_effect") or "",
            "powerPointAnimationClass": item.get("animation_class") or "",
            "transitionName": DEFAULT_TRANSITION_NAME,
        },
        "animationTracks": {},
    }


def callout_clip(clip_id: int, item: dict[str, Any], slide_duration: int) -> dict[str, Any]:
    text = safe_text(item.get("text"))
    width = max(float(item.get("width_px") or 300), 20.0)
    height = max(float(item.get("height_px") or 50), 16.0)
    x = float(item.get("x_px") or 0.0)
    y = float(item.get("y_px") or 0.0)
    font_size = STANDARD_FONT_SIZE_PT
    font_name = STANDARD_FONT_NAME
    red, green, blue = hex_to_float_rgb(item.get("font_color"))
    color = camtasia_color(item.get("font_color"))
    weight_value = 700 if item.get("bold") else 400
    weight_name = "Bold" if item.get("bold") else "Regular"
    italic = 1 if item.get("italic") else 0
    underline = 1 if item.get("underline") else 0
    align = item.get("alignment") or "left"
    range_end = len(text)
    translation0 = x + width / 2.0 - CANVAS_WIDTH / 2.0
    translation1 = -(y + height / 2.0 - CANVAS_HEIGHT / 2.0)
    text_attrs = [
        {"name": "fontSize", "rangeEnd": range_end, "rangeStart": 0, "value": font_size, "valueType": "double"},
        {"name": "fontName", "rangeEnd": range_end, "rangeStart": 0, "value": font_name, "valueType": "string"},
        {"name": "strikethrough", "rangeEnd": range_end, "rangeStart": 0, "value": 0, "valueType": "int"},
        {"name": "kerning", "rangeEnd": range_end, "rangeStart": 0, "value": 0.0, "valueType": "double"},
        {"name": "fgColor", "rangeEnd": range_end, "rangeStart": 0, "value": color, "valueType": "color"},
        {"name": "fontWeight", "rangeEnd": range_end, "rangeStart": 0, "value": weight_value, "valueType": "int"},
        {"name": "underline", "rangeEnd": range_end, "rangeStart": 0, "value": underline, "valueType": "int"},
        {"name": "fontItalic", "rangeEnd": range_end, "rangeStart": 0, "value": italic, "valueType": "int"},
    ]
    start_ticks, duration_ticks = item_clip_timing(item, slide_duration)
    return {
        "id": clip_id,
        "_type": "Callout",
        "attributes": {"ident": item["id"], "autoRotateText": True},
        "def": {
            "kind": "remix",
            "shape": "text",
            "style": "basic",
            "enable-ligatures": 1.0,
            "height": height,
            "line-spacing": 0.0,
            "width": width,
            "word-wrap": 1.0,
            "horizontal-alignment": align,
            "resize-behavior": "resizeText",
            "text": text,
            "vertical-alignment": "center",
            "font": {
                "color-blue": blue,
                "color-green": green,
                "color-red": red,
                "color-opacity": 1.0,
                "size": font_size,
                "tracking": 0.0,
                "name": font_name,
                "weight": weight_name,
            },
            "textAttributes": {
                "type": "textAttributeList",
                "keyframes": [{"endTime": 0, "time": 0, "value": text_attrs, "duration": 0}],
            },
        },
        "parameters": {
            "translation0": translation0,
            "translation1": translation1,
            "scale0": 1.0,
            "scale1": 1.0,
            "geometryCrop0": 0.0,
            "geometryCrop1": 0.0,
            "geometryCrop2": 0.0,
            "geometryCrop3": 0.0,
        },
        "effects": [],
        "start": start_ticks,
        "duration": duration_ticks,
        "mediaStart": 0,
        "mediaDuration": duration_ticks,
        "scalar": 1,
        "metadata": {
            "audiateLinkedSession": "",
            "clipSpeedAttribute": {"type": "bool", "value": False},
            "default-height": {"type": "double", "value": height},
            "default-width": {"type": "double", "value": width},
            "default-scale0": {"type": "double", "value": 1.0},
            "default-scale1": {"type": "double", "value": 1.0},
            "default-translation0": {"type": "double", "value": translation0},
            "default-translation1": {"type": "double", "value": translation1},
            "powerPointElementType": "editable_text",
            "powerPointShapeId": {"type": "int", "value": int(item.get("shape_id") or 0)},
            "powerPointShapeName": item.get("shape_name") or "",
            "powerPointReadingOrder": {"type": "int", "value": int(item.get("reading_order") or 0)},
            "powerPointRelationshipId": item.get("relationship_id") or "",
            "powerPointRelationshipReason": item.get("relationship_reason") or "",
            "powerPointRotation": {"type": "double", "value": float(item.get("rotation") or 0.0)},
            "powerPointZOrderPath": ".".join(str(v) for v in item.get("z_order_path") or []),
            "powerPointAnimationOrder": {"type": "int", "value": int(item.get("animation_order") or 0)},
            "powerPointAnimationEffect": item.get("animation_effect") or "",
            "powerPointAnimationClass": item.get("animation_class") or "",
            "transitionName": DEFAULT_TRANSITION_NAME,
        },
        "animationTracks": {},
    }


def slide_group(group_id: int, slide: dict[str, Any], tracks: list[dict[str, Any]]) -> dict[str, Any]:
    start = sec_to_ticks(slide["start_time_sec"])
    duration = sec_to_ticks(slide["duration_sec"])
    name = f"Slide {int(slide['slide_number']):02d} - {clean_name(slide.get('title') or '')}"
    return {
        "id": group_id,
        "_type": "Group",
        "start": start,
        "duration": duration,
        "mediaStart": 0,
        "mediaDuration": duration,
        "scalar": 1,
        "metadata": {
            "audiateLinkedSession": "",
            "clipSpeedAttribute": {"type": "bool", "value": False},
            "colorAttribute": {"type": "color", "value": [0, 0, 0, 0]},
            "effectApplied": "none",
            "isOpen": {"type": "bool", "value": False},
            "powerPointElementType": "slide_group",
        },
        "animationTracks": {},
        "parameters": {
            "translation0": 0.0,
            "translation1": 0.0,
            "scale0": 1.0,
            "scale1": 1.0,
            "geometryCrop0": 0.0,
            "geometryCrop1": 0.0,
            "geometryCrop2": 0.0,
            "geometryCrop3": 0.0,
            "volume": 1.0,
        },
        "effects": [],
        "attributes": {
            "ident": name,
            "widthAttr": float(CANVAS_WIDTH),
            "heightAttr": float(CANVAS_HEIGHT),
            "maxDurationAttr": 0,
            "gain": 1.0,
            "mixToMono": False,
            "assetProperties": [],
        },
        "tracks": tracks,
    }



def _visual_bounds(item: dict[str, Any]) -> tuple[float, float, float, float]:
    x = float(item.get("x_px") or 0.0)
    y = float(item.get("y_px") or 0.0)
    w = max(float(item.get("width_px") or 0.0), 0.01)
    h = max(float(item.get("height_px") or 0.0), 0.01)
    return x, y, x + w, y + h


def _bbox_union(items: list[dict[str, Any]]) -> tuple[float, float, float, float]:
    bounds = [_visual_bounds(item) for item in items]
    return (
        min(b[0] for b in bounds),
        min(b[1] for b in bounds),
        max(b[2] for b in bounds),
        max(b[3] for b in bounds),
    )


def _bbox_area(bounds: tuple[float, float, float, float]) -> float:
    return max(bounds[2] - bounds[0], 0.0) * max(bounds[3] - bounds[1], 0.0)


def _contains_item(container: dict[str, Any], child: dict[str, Any]) -> bool:
    if container is child:
        return False
    cx1, cy1, cx2, cy2 = _visual_bounds(container)
    x1, y1, x2, y2 = _visual_bounds(child)
    center_x = (x1 + x2) / 2.0
    center_y = (y1 + y2) / 2.0
    if not (cx1 - 2.0 <= center_x <= cx2 + 2.0 and cy1 - 2.0 <= center_y <= cy2 + 2.0):
        return False
    overlap_w = max(0.0, min(cx2, x2) - max(cx1, x1))
    overlap_h = max(0.0, min(cy2, y2) - max(cy1, y1))
    child_area = max((x2 - x1) * (y2 - y1), 0.01)
    return (overlap_w * overlap_h) / child_area >= 0.60


def _is_structural_visual(item: dict[str, Any]) -> bool:
    if item.get("element_kind") != "shape":
        return False
    x1, y1, x2, y2 = _visual_bounds(item)
    width_ratio = max(x2 - x1, 0.0) / float(CANVAS_WIDTH)
    height_ratio = max(y2 - y1, 0.0) / float(CANVAS_HEIGHT)
    area_ratio = _bbox_area((x1, y1, x2, y2)) / float(CANVAS_WIDTH * CANVAS_HEIGHT)
    return (
        area_ratio >= 0.40
        or (width_ratio >= 0.88 and height_ratio >= 0.15)
        or (height_ratio >= 0.88 and width_ratio >= 0.08)
    )


def _z_key(item: dict[str, Any]) -> tuple[int, ...]:
    raw = item.get("z_order_path") or []
    try:
        return tuple(int(value) for value in raw)
    except Exception:
        return (int(item.get("shape_id") or 0),)


def _spatial_rows(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if not records:
        return []
    pending = sorted(records, key=lambda r: (float(r["top"]), float(r["left"]), int(r.get("stable") or 0)))
    rows: list[list[dict[str, Any]]] = []
    row_anchor_tops: list[float] = []
    for record in pending:
        top = float(record["top"])
        if rows and abs(top - row_anchor_tops[-1]) <= 24.0:
            rows[-1].append(record)
            row_anchor_tops[-1] = min(row_anchor_tops[-1], top)
        else:
            rows.append([record])
            row_anchor_tops.append(top)
    ordered: list[dict[str, Any]] = []
    for row in rows:
        row.sort(key=lambda r: (float(r["left"]), float(r["top"]), int(r.get("stable") or 0)))
        ordered.extend(row)
    return ordered


def spatial_order_for_slide(extracted: ExtractedData, slide_number: int) -> list[dict[str, Any]]:
    """Return one mixed shape/text sequence using slide position and visual relationships.

    Existing shape and text extraction remains intact. This function only adds the
    visual sequence used by the Camtasia timeline and the combined extraction manifest.
    """
    items: list[dict[str, Any]] = []
    for item in extracted.shape_items:
        if int(item.get("slide_number") or 0) == slide_number:
            item["element_kind"] = "shape"
            items.append(item)
    for item in extracted.text_items:
        if int(item.get("slide_number") or 0) == slide_number:
            item["element_kind"] = "text"
            items.append(item)
    if not items:
        return []

    parent = list(range(len(items)))
    reasons: dict[int, set[str]] = {i: set() for i in range(len(items))}

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def union(a: int, b: int, reason: str) -> None:
        ra, rb = find(a), find(b)
        if ra == rb:
            reasons[ra].add(reason)
            return
        if ra > rb:
            ra, rb = rb, ra
        parent[rb] = ra
        reasons[ra].update(reasons.pop(rb, set()))
        reasons[ra].add(reason)

    # A PowerPoint shape and its editable text representation are one logical object.
    by_shape_id: dict[int, list[int]] = {}
    for index, item in enumerate(items):
        by_shape_id.setdefault(int(item.get("shape_id") or 0), []).append(index)
    for members in by_shape_id.values():
        if len(members) > 1:
            first = members[0]
            for member in members[1:]:
                union(first, member, "same_powerpoint_shape")

    # Keep explicit PowerPoint group children related without grouping unrelated top-level items.
    by_parent_group: dict[str, list[int]] = {}
    for index, item in enumerate(items):
        group_path = str(item.get("parent_group_path") or "")
        if group_path and not _is_structural_visual(item):
            by_parent_group.setdefault(group_path, []).append(index)
    for members in by_parent_group.values():
        if 1 < len(members) <= 12:
            first = members[0]
            for member in members[1:]:
                union(first, member, "powerpoint_group")

    # Associate text/icons/pictures with the smallest reasonable card/container shape.
    container_indices = [
        i for i, item in enumerate(items)
        if item.get("element_kind") == "shape"
        and item.get("asset_kind") != "picture"
        and not _is_structural_visual(item)
        and _bbox_area(_visual_bounds(item)) <= CANVAS_WIDTH * CANVAS_HEIGHT * 0.38
    ]
    for child_index, child in enumerate(items):
        candidates: list[tuple[float, int]] = []
        child_area = _bbox_area(_visual_bounds(child))
        for container_index in container_indices:
            if container_index == child_index:
                continue
            container = items[container_index]
            container_area = _bbox_area(_visual_bounds(container))
            if container_area <= child_area * 1.02:
                continue
            if _contains_item(container, child):
                candidates.append((container_area, container_index))
        if candidates:
            _, best = min(candidates)
            union(best, child_index, "spatial_container")

    clusters: dict[int, list[int]] = {}
    for index in range(len(items)):
        clusters.setdefault(find(index), []).append(index)

    structural_clusters: list[dict[str, Any]] = []
    content_clusters: list[dict[str, Any]] = []
    for stable, (root, member_indices) in enumerate(clusters.items()):
        members = [items[i] for i in member_indices]
        bounds = _bbox_union(members)
        record = {
            "root": root,
            "members": members,
            "bounds": bounds,
            "left": bounds[0],
            "top": bounds[1],
            "stable": stable,
            "reasons": sorted(reasons.get(find(root), set())),
        }
        if all(_is_structural_visual(member) for member in members):
            structural_clusters.append(record)
        else:
            content_clusters.append(record)

    # Structural/background-like slide objects stay at the bottom and do not pull
    # unrelated content into one large relationship group.
    structural_clusters.sort(key=lambda r: (_z_key(r["members"][0]), r["top"], r["left"]))
    ordered_clusters = structural_clusters + _spatial_rows(content_clusters)

    ordered: list[dict[str, Any]] = []
    relationship_number = 0
    for cluster in ordered_clusters:
        relationship_number += 1
        members = list(cluster["members"])
        relationship_id = f"S{slide_number:03d}_R{relationship_number:03d}"

        # Put only true outer card/container shapes first. Nested cards and their
        # text/icons remain in the spatial sequence so the result becomes
        # Shape -> related Text/Icon -> next Shape, not all Shapes -> all Text.
        container_candidates: list[dict[str, Any]] = []
        for member in members:
            if member.get("element_kind") != "shape":
                continue
            contained_count = sum(1 for other in members if other is not member and _contains_item(member, other))
            if contained_count >= 2:
                container_candidates.append(member)

        outer_containers: list[dict[str, Any]] = []
        for candidate in container_candidates:
            if any(
                other is not candidate
                and _bbox_area(_visual_bounds(other)) > _bbox_area(_visual_bounds(candidate))
                and _contains_item(other, candidate)
                for other in container_candidates
            ):
                continue
            outer_containers.append(candidate)

        outer_ids = {id(item) for item in outer_containers}
        content = [member for member in members if id(member) not in outer_ids]
        outer_containers.sort(
            key=lambda item: (_z_key(item), float(item.get("y_px") or 0.0), float(item.get("x_px") or 0.0))
        )
        content_records = []
        for stable, member in enumerate(content):
            content_records.append({
                "item": member,
                "left": float(member.get("x_px") or 0.0),
                "top": float(member.get("y_px") or 0.0),
                "stable": stable,
            })
        content_order = [record["item"] for record in _spatial_rows(content_records)]

        cluster_order = outer_containers + content_order
        for member in cluster_order:
            member["relationship_id"] = relationship_id
            member["relationship_size"] = len(members)
            member["relationship_reason"] = ",".join(cluster.get("reasons") or []) or "standalone"
            ordered.append(member)

    for reading_index, item in enumerate(ordered, start=1):
        item["reading_order"] = reading_index
    return ordered


def visual_manifest(extracted: ExtractedData) -> list[dict[str, Any]]:
    manifest: list[dict[str, Any]] = []
    for slide_number in range(1, extracted.slide_count + 1):
        for item in spatial_order_for_slide(extracted, slide_number):
            manifest.append({
                "id": item.get("id"),
                "slide_number": slide_number,
                "reading_order": item.get("reading_order"),
                "relationship_id": item.get("relationship_id"),
                "relationship_reason": item.get("relationship_reason"),
                "element_kind": item.get("element_kind"),
                "asset_kind": item.get("asset_kind"),
                "shape_id": item.get("shape_id"),
                "shape_name": item.get("shape_name"),
                "shape_path": item.get("shape_path"),
                "text": item.get("text", ""),
                "x_px": item.get("x_px"),
                "y_px": item.get("y_px"),
                "width_px": item.get("width_px"),
                "height_px": item.get("height_px"),
                "x_emu": item.get("x_emu"),
                "y_emu": item.get("y_emu"),
                "width_emu": item.get("width_emu"),
                "height_emu": item.get("height_emu"),
                "rotation": item.get("rotation"),
                "alignment": item.get("alignment"),
                "z_order_path": item.get("z_order_path"),
                "group_depth": item.get("group_depth"),
                "parent_group_path": item.get("parent_group_path"),
            })
    return manifest

def build_project(
    extracted: ExtractedData,
    backgrounds: list[Path],
    shape_assets: dict[str, str],
    project_root: Path,
) -> dict[str, Any]:
    source_bin: list[dict[str, Any]] = []
    slide_groups: list[dict[str, Any]] = []
    next_source_id = 1
    next_clip_id = 100

    for slide in extracted.slide_timing:
        slide_number = int(slide["slide_number"])
        background = backgrounds[slide_number - 1]
        source_id = next_source_id
        next_source_id += 1
        source_bin.append(make_source(source_id, background, project_root))
        duration_ticks = sec_to_ticks(float(slide["duration_sec"]))
        bg_clip = background_clip(next_clip_id, source_id, duration_ticks, slide_number)
        internal_tracks = [
            track_wrapper(
                0,
                f"S{slide_number:03d} | 01 BACKGROUND | Clean PowerPoint background",
                [bg_clip],
                fade_transitions_for_clip(bg_clip),
            )
        ]
        next_clip_id += 1
        track_index = 1

        # IMPORTANT: use the mixed spatial/relationship sequence. Do not put every
        # shape first and every text item afterwards. Each visual object remains a
        # separate editable/movable Camtasia layer exactly as before.
        for item in spatial_order_for_slide(extracted, slide_number):
            element_kind = item.get("element_kind")
            if element_kind == "shape":
                asset = shape_assets.get(item["id"])
                if not asset:
                    continue
                asset_path = Path(asset)
                source_id = next_source_id
                next_source_id += 1
                source_bin.append(make_source(source_id, asset_path, project_root))
                layer_kind = "IMAGE PNG" if item.get("asset_kind") == "picture" else "SHAPE PNG"
                label = (
                    f"S{slide_number:03d} | {track_index + 1:02d} {layer_kind} | "
                    f"{clean_name(item.get('shape_name', '')[:50])}"
                )
                clip = image_shape_clip(next_clip_id, source_id, item, asset_path, duration_ticks)
            else:
                label = (
                    f"S{slide_number:03d} | {track_index + 1:02d} EDITABLE TEXT | "
                    f"{clean_name(item.get('text', '')[:50])}"
                )
                clip = callout_clip(next_clip_id, item, duration_ticks)

            internal_tracks.append(track_wrapper(track_index, label, [clip], fade_transitions_for_clip(clip)))
            next_clip_id += 1
            track_index += 1

        slide_groups.append(slide_group(next_clip_id, slide, internal_tracks))
        next_clip_id += 1

    total_duration = sec_to_ticks(extracted.slide_timing[-1]["end_time_sec"] if extracted.slide_timing else 0)
    return {
        "title": "PowerPoint to editable Camtasia - verified object timeline",
        "description": "Clean backgrounds, imported visual layers, Open Sans 34 editable text, Fade transitions, and grouped slide timelines.",
        "author": "",
        "targetLoudness": -18.0,
        "shouldApplyLoudnessNormalization": True,
        "videoFormatFrameRate": 30,
        "audioFormatSampleRate": 44100,
        "allowSubFrameEditing": False,
        "width": float(CANVAS_WIDTH),
        "height": float(CANVAS_HEIGHT),
        "version": "9.0",
        "editRate": TICKS_PER_SECOND,
        "authoringClientName": {"name": "Camtasia", "platform": "Windows", "version": "2025.2.5"},
        "sourceBin": source_bin,
        "timeline": {
            "id": 2,
            "sceneTrack": {
                "scenes": [{
                    "duration": total_duration,
                    "title": "",
                    "type": "",
                    "csml": {
                        "tracks": [track_wrapper(0, "SLIDE GROUPS - proper timeline", slide_groups)]
                    },
                }]
            },
        },
        "metadata": {
            "Date": datetime.now().strftime("%Y-%m-%d %I:%M:%S %p"),
            "IsStandalone": True,
            "Language": "ENU",
            "ProjectDimensionsChanged": "1",
            "Title": "PowerPoint to editable Camtasia - verified object timeline",
            "RulerGuidesX": [],
            "RulerGuidesY": [],
            "RulerShowing": False,
        },
    }

def write_reports(
    project: dict[str, Any],
    extracted: ExtractedData,
    project_root: Path,
    backgrounds: list[Path],
    shape_assets: dict[str, str],
    background_mode: str,
) -> None:
    reports = ensure_dir(project_root / "reports")
    slide_rows = []
    track_rows = []
    edit_rows = []
    animation_rows = []
    expected_visual_items = [
        item for item in extracted.shape_items
        if item.get("classification") == "separate_png_layer"
    ]
    missing_visual_items = [
        item for item in expected_visual_items
        if item.get("id") not in shape_assets
    ]
    for event in extracted.animation_items:
        animation_rows.append({
            "slide": event.get("slide_number"),
            "shape_id": event.get("shape_id"),
            "order": event.get("order"),
            "class": event.get("animation_class"),
            "action": event.get("action"),
            "effect": event.get("effect"),
            "start_sec": event.get("start_time_sec"),
            "effect_duration_sec": event.get("effect_duration_sec"),
        })
    for slide in extracted.slide_timing:
        slide_number = int(slide["slide_number"])
        items = [item for item in extracted.text_items if int(item["slide_number"]) == slide_number]
        shapes = [item for item in extracted.shape_items if int(item["slide_number"]) == slide_number]
        slide_rows.append({
            "slide": slide_number,
            "title": slide.get("title", ""),
            "start_sec": slide["start_time_sec"],
            "duration_sec": slide["duration_sec"],
            "end_sec": slide["end_time_sec"],
            "editable_text_count": len(items),
            "shape_count": len(shapes),
            "native_shape_candidate_count": sum(1 for item in shapes if item.get("editable_candidate") == "native_camtasia_shape"),
            "separate_shape_png_count": sum(1 for item in shapes if item.get("classification") == "separate_png_layer" and item["id"] in shape_assets),
            "background": backgrounds[slide_number - 1].name,
            "transition_name": DEFAULT_TRANSITION_NAME,
            "font_name": STANDARD_FONT_NAME,
            "font_size_pt": STANDARD_FONT_SIZE_PT,
        })
        track_rows.append({
            "slide": slide_number,
            "reading_order": 0,
            "relationship_id": "",
            "track": f"S{slide_number:03d} | 01 BACKGROUND | Clean PowerPoint background",
            "type": "background",
            "editable": False,
            "visible": True,
            "transition_name": DEFAULT_TRANSITION_NAME,
        })
        visual_items = spatial_order_for_slide(extracted, slide_number)
        for index, item in enumerate(visual_items, start=1):
            if item.get("element_kind") == "shape":
                if item["id"] in shape_assets:
                    track_type = "separate_image_png" if item.get("asset_kind") == "picture" else "separate_shape_png"
                    visible = True
                else:
                    track_type = "shape_export_missing"
                    visible = False
                editable = False
                label_text = clean_name(item.get("shape_name", "")[:50])
                classification = "Separate image PNG layer" if item.get("asset_kind") == "picture" else "Separate shape PNG layer"
                text_value = ""
            else:
                track_type = "editable_text"
                visible = True
                editable = True
                label_text = clean_name(item.get("text", "")[:50])
                classification = "Native editable Camtasia text"
                text_value = item.get("text")

            track_rows.append({
                "slide": slide_number,
                "reading_order": item.get("reading_order"),
                "relationship_id": item.get("relationship_id"),
                "relationship_reason": item.get("relationship_reason"),
                "track": f"S{slide_number:03d} | {index + 1:02d} {track_type.upper()} | {label_text}",
                "type": track_type,
                "editable": editable,
                "visible": visible,
                "start_sec": item.get("start_time_sec"),
                "duration_sec": item.get("duration_sec"),
                "animation_order": item.get("animation_order"),
                "animation_effect": item.get("animation_effect"),
                "transition_name": DEFAULT_TRANSITION_NAME,
                "font_name": STANDARD_FONT_NAME if editable else "",
                "font_size_pt": STANDARD_FONT_SIZE_PT if editable else "",
            })
            edit_rows.append({
                "slide": slide_number,
                "reading_order": item.get("reading_order"),
                "relationship_id": item.get("relationship_id"),
                "relationship_reason": item.get("relationship_reason"),
                "element_kind": item.get("element_kind"),
                "shape_id": item.get("shape_id"),
                "text": text_value,
                "classification": classification,
                "x_px": item.get("x_px"),
                "y_px": item.get("y_px"),
                "width_px": item.get("width_px"),
                "height_px": item.get("height_px"),
                "x_emu": item.get("x_emu"),
                "y_emu": item.get("y_emu"),
                "width_emu": item.get("width_emu"),
                "height_emu": item.get("height_emu"),
                "rotation": item.get("rotation"),
                "alignment": item.get("alignment"),
                "z_order_path": ".".join(str(v) for v in item.get("z_order_path") or []),
                "group_depth": item.get("group_depth"),
                "parent_group_path": item.get("parent_group_path"),
                "start_sec": item.get("start_time_sec"),
                "duration_sec": item.get("duration_sec"),
                "transition_name": DEFAULT_TRANSITION_NAME,
                "editable_candidate": item.get("editable_candidate"),
                "font_name": STANDARD_FONT_NAME if editable else "",
                "font_size_pt": STANDARD_FONT_SIZE_PT if editable else "",
            })

    def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
        if not rows:
            path.write_text("", encoding="utf-8")
            return
        fieldnames: list[str] = []
        for row in rows:
            for key in row.keys():
                if key not in fieldnames:
                    fieldnames.append(key)
        with path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)

    write_csv(reports / "slide_map.csv", slide_rows)
    write_csv(reports / "timing_report.csv", slide_rows)
    write_csv(reports / "track_map.csv", track_rows)
    write_csv(reports / "animation_map.csv", animation_rows)
    write_csv(reports / "editability_report.csv", edit_rows)
    missing_lines: list[str] = []
    missing_backgrounds = [path for path in backgrounds if not path.exists() or path.stat().st_size <= 0]
    if missing_backgrounds:
        missing_lines.append("FAIL: missing or empty background PNG files:")
        missing_lines.extend(f"- {path.name}" for path in missing_backgrounds)
    if missing_visual_items:
        missing_lines.append("FAIL: expected visual PNG layers were not exported/imported:")
        for item in missing_visual_items:
            missing_lines.append(
                f"- S{int(item.get('slide_number') or 0):03d} "
                f"shape_id={item.get('shape_id')} name={item.get('shape_name')} path={item.get('shape_path')}"
            )
    if not missing_lines:
        missing_lines.append("PASS: all background PNG files and expected visual PNG layers are present.")
    (reports / "missing_assets_report.txt").write_text("\n".join(missing_lines) + "\n", encoding="utf-8")

    numeric_refs: list[int] = []
    clip_ids: list[int] = []

    def walk(obj: Any) -> None:
        if isinstance(obj, dict):
            if obj.get("_type") and isinstance(obj.get("id"), int):
                clip_ids.append(obj["id"])
            if isinstance(obj.get("src"), int):
                numeric_refs.append(obj["src"])
            for value in obj.values():
                walk(value)
        elif isinstance(obj, list):
            for value in obj:
                walk(value)

    walk(project.get("timeline", {}))
    source_ids = {item["id"] for item in project.get("sourceBin", [])}
    errors = []
    missing = sorted(set(numeric_refs) - source_ids)
    if missing:
        errors.append({"check": "missing_source_refs", "values": missing})
    duplicates = [clip_id for clip_id, count in __import__("collections").Counter(clip_ids).items() if count > 1]
    if duplicates:
        errors.append({"check": "duplicate_clip_ids", "values": duplicates})
    if len(backgrounds) != extracted.slide_count:
        errors.append({"check": "background_count", "expected": extracted.slide_count, "actual": len(backgrounds)})
    if missing_visual_items:
        errors.append({
            "check": "missing_visual_png_layers",
            "expected": len(expected_visual_items),
            "actual": len(shape_assets),
            "missing": [
                {
                    "slide_number": item.get("slide_number"),
                    "shape_id": item.get("shape_id"),
                    "shape_name": item.get("shape_name"),
                    "shape_path": item.get("shape_path"),
                }
                for item in missing_visual_items
            ],
        })
    qc = {
        "status": "PASS" if not errors else "FAIL",
        "checkedOn": datetime.now().isoformat(timespec="seconds"),
        "slideCount": extracted.slide_count,
        "sourceCount": len(project.get("sourceBin", [])),
        "editableTextCount": len(extracted.text_items),
        "shapeCount": len(extracted.shape_items),
        "nativeShapeCandidateCount": sum(1 for item in extracted.shape_items if item.get("editable_candidate") == "native_camtasia_shape"),
        "expectedVisualPngLayerCount": len(expected_visual_items),
        "importedVisualPngLayerCount": len(shape_assets),
        "separateShapePngCount": len(shape_assets),
        "separateImagePngCount": sum(1 for item in extracted.shape_items if item.get("asset_kind") == "picture" and item["id"] in shape_assets),
        "animationEventCount": len(extracted.animation_items),
        "transitionName": DEFAULT_TRANSITION_NAME,
        "editableTextFont": STANDARD_FONT_NAME,
        "editableTextSizePt": STANDARD_FONT_SIZE_PT,
        "backgroundMode": background_mode,
        "errors": errors,
        "notes": [
            "This runner creates clean slide/master backgrounds, separated visual PNG layers, and editable Camtasia text.",
            "All visible non-text PowerPoint shapes are exported as individual PNG layers so PowerPoint-only shapes do not disappear in Camtasia.",
            "Pictures, people, bars, cards, borders, icons, and complex shapes can be moved, hidden, replaced, or deleted as separate timeline layers.",
            "PowerPoint animation order and delays are mapped to Camtasia clip start/duration when timing XML targets can be read.",
            "If backgroundMode is no_text_fallback, PowerPoint rejected the background-only PPTX and backgrounds may include non-text shapes.",
            "Complex visual effects such as Morph, curved motion paths, and advanced triggers are reported for manual Camtasia refinement.",
        ],
    }
    save_json(qc, reports / "qc_report.json")


def write_guides(project_root: Path, project_file_name: str = "project.tscproj") -> None:
    (project_root / "OPEN_THIS_PROJECT.txt").write_text(
        f"Open {project_file_name} in Camtasia.\n"
        "Each slide is one expandable group. Tracks are labelled IMAGE PNG, SHAPE PNG, or EDITABLE TEXT.\n",
        encoding="utf-8",
    )
    (project_root / "README_EDITING_GUIDE.txt").write_text(
        "README EDITING GUIDE\n\n"
        f"1. Open {project_file_name} in Camtasia.\n"
        "2. Expand a Slide group on the timeline.\n"
        "3. Edit tracks labelled EDITABLE TEXT directly in Camtasia. Editable text is created as Open Sans 34 pt.\n"
        "4. Move, hide, replace, or delete tracks labelled IMAGE PNG or SHAPE PNG as separate visual layers.\n"
        "5. Colored bars, rounded cards, borders, pictures, and icons are imported as PNG layers for accurate display.\n"
        "6. Background tracks are base slide/master PNGs only. Imported timeline items use Fade transitions.\n"
        "7. Check reports/animation_map.csv for the mapped PowerPoint animation order, delay, effect, and target shape.\n"
        "8. If PowerPoint could not export clean backgrounds, backgrounds may include shapes too. "
        "Check reports/qc_report.json for backgroundMode. If duplicated, hide or delete the extra shape track.\n"
        "9. Reports are in the reports folder.\n",
        encoding="utf-8",
    )
    (project_root / "KNOWN_LIMITATIONS.txt").write_text(
        "KNOWN LIMITATIONS\n\n"
        "- This runner prioritizes proper timeline structure, editable text, and object-level editability.\n"
        "- All visible PowerPoint shapes are imported as separate PNG layers for reliable visual accuracy in Camtasia.\n"
        "- People/photos/pictures, colored bars, rounded boxes, borders, and complex artwork can be moved/replaced/deleted but not color-edited like native Camtasia shapes.\n"
        "- Backgrounds are exported from PowerPoint after slide-level shapes are removed, leaving the slide/master background clean where PowerPoint allows it.\n"
        "- Editable text is standardized to Open Sans 34 pt, and imported timeline items use Fade transitions.\n"
        "- The runner fails closed and does not package the final ZIP if any expected visual PNG layer is missing.\n"
        "- PowerPoint animation order and delays are mapped to clip timing. Exact native recreation of Morph, complex motion paths, 3D rotation, and advanced triggers may still need manual Camtasia refinement.\n",
        encoding="utf-8",
    )


def zip_project(project_root: Path, output_zip: Path) -> None:
    if output_zip.exists():
        output_zip.unlink()
    with zipfile.ZipFile(output_zip, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for file in project_root.rglob("*"):
            if file.is_file():
                z.write(file, file.relative_to(project_root.parent).as_posix())
    with zipfile.ZipFile(output_zip) as z:
        bad = z.testzip()
    if bad:
        die("Created ZIP failed integrity check at: " + bad)


def write_same_location_project_copy(
    project: dict[str, Any],
    output_parent: Path,
    project_folder_name: str,
    project_file_name: str,
) -> Path:
    """Write a convenience .tscproj beside the selected PPT without breaking media links.

    The canonical self-contained project keeps its existing ./media/... references.
    This same-location copy points those references into the companion project folder.
    Timeline data, editability, animation timing, tracks, and all other project content
    remain unchanged.
    """
    sibling_project = json.loads(json.dumps(project))
    media_prefix = f"./{project_folder_name}/"
    for source in sibling_project.get("sourceBin", []):
        src = source.get("src")
        if isinstance(src, str) and src.startswith("./"):
            source["src"] = media_prefix + src[2:]
    destination = output_parent / project_file_name
    destination.write_text(json.dumps(sibling_project, indent=2, ensure_ascii=False), encoding="utf-8")
    return destination


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert a PowerPoint presentation into the existing editable Camtasia project format."
    )
    parser.add_argument(
        "--source",
        type=str,
        help="PowerPoint .pptx file to convert. If omitted, legacy input-folder mode is used.",
    )
    parser.add_argument(
        "--output-parent",
        type=str,
        help="Folder where the Camtasia project folder and ZIP are created. Defaults to the selected PPT folder.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    root = Path(__file__).resolve().parents[1]
    input_dir = root / "input"
    working_dir = root / "working"

    # UI/CLI mode: use the explicitly selected PPT and save the finished output
    # beside it. Legacy mode remains available when --source is not supplied.
    if args.source:
        source_pptx = Path(args.source).expanduser().resolve()
        if not source_pptx.exists() or not source_pptx.is_file():
            die(f"Selected PowerPoint file does not exist: {source_pptx}")
        if source_pptx.suffix.lower() != ".pptx":
            die(f"Selected file is not a .pptx file: {source_pptx}")
        if source_pptx.name.startswith("~$"):
            die("Please select the actual PowerPoint file, not a temporary ~$ file.")

        output_parent = (
            Path(args.output_parent).expanduser().resolve()
            if args.output_parent
            else source_pptx.parent
        )
        ensure_dir(output_parent)
        safe_stem = source_pptx.stem
        project_folder_name = f"{safe_stem}_CAMTASIA_PROJECT"
        project_file_name = f"{safe_stem}.tscproj"
        project_root = output_parent / project_folder_name
        output_zip = output_parent / f"{project_folder_name}.zip"
        direct_project_file = output_parent / project_file_name

        if project_root.exists():
            shutil.rmtree(project_root)
        if output_zip.exists():
            output_zip.unlink()
        if direct_project_file.exists():
            direct_project_file.unlink()
    else:
        # Original behaviour retained for backwards compatibility.
        output_dir = root / "output"
        project_root = output_dir / "FINAL_CAMTASIA_PROJECT"
        output_zip = output_dir / "FINAL_CAMTASIA_PROJECT.zip"
        project_file_name = "project.tscproj"
        direct_project_file = None
        ensure_dir(input_dir)
        if output_dir.exists():
            shutil.rmtree(output_dir)
        source_pptx = find_input_pptx(input_dir)

    media_backgrounds = project_root / "media" / "backgrounds"
    media_images = project_root / "media" / "images"
    reports_dir = project_root / "reports"

    if working_dir.exists():
        shutil.rmtree(working_dir)
    ensure_dir(working_dir)
    ensure_dir(media_backgrounds)
    ensure_dir(media_images)
    ensure_dir(reports_dir)

    print("Using PowerPoint:")
    print(source_pptx)
    print("Output project folder:")
    print(project_root)
    print("Camtasia project file name:")
    print(project_file_name)

    extracted = extract_text_and_timing(source_pptx, working_dir)
    no_text_pptx = working_dir / "source_no_text.pptx"
    background_only_pptx = working_dir / "source_background_only.pptx"
    duplicate_without_text(source_pptx, no_text_pptx, working_dir / "remove_text_report.json")
    shape_assets = export_shape_pngs_with_powerpoint(no_text_pptx, media_images, extracted.shape_items)
    backgrounds, background_mode = export_backgrounds_safely(
        source_pptx,
        no_text_pptx,
        background_only_pptx,
        media_backgrounds,
        working_dir,
        extracted.slide_count,
    )
    project = build_project(extracted, backgrounds, shape_assets, project_root)
    project_file = project_root / project_file_name
    project_file.write_text(json.dumps(project, indent=2, ensure_ascii=False), encoding="utf-8")
    write_reports(project, extracted, project_root, backgrounds, shape_assets, background_mode)
    qc_path = project_root / "reports" / "qc_report.json"
    qc = json.loads(qc_path.read_text(encoding="utf-8"))
    if qc.get("status") != "PASS":
        print("")
        print("QC FAILED: the project was created for inspection, but it was not packaged.")
        print("Open this report:")
        print(qc_path)
        die("One or more required visual layers are missing or invalid. Check reports\\missing_assets_report.txt.")
    write_guides(project_root, project_file_name)
    shutil.copy2(source_pptx, project_root / ("SOURCE_" + source_pptx.name))
    zip_project(project_root, output_zip)
    if args.source:
        direct_project_file = write_same_location_project_copy(
            project,
            source_pptx.parent if not args.output_parent else Path(args.output_parent).expanduser().resolve(),
            project_root.name,
            project_file_name,
        )

    print("")
    print("DONE")
    print("Project folder:")
    print(project_root)
    print("Project file (self-contained folder):")
    print(project_file)
    if direct_project_file is not None:
        print("Same-location Camtasia project file:")
        print(direct_project_file)
    print("Final ZIP:")
    print(output_zip)


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        die("Stopped by user.")
